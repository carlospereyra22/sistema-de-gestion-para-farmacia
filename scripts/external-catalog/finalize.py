"""Generate deliverables from a successfully committed local import, not a download estimate."""
import csv
import hashlib
import json
import pathlib
import shutil
import sys
from contextlib import ExitStack
from collections import Counter, defaultdict
from openpyxl import load_workbook

prepared, target = map(pathlib.Path, sys.argv[1:3])
result = json.loads((prepared / "postgres-result.json").read_text(encoding="utf-8"))
if result["status"] != "IMPORTED_TO_LOCAL_POSTGRES": raise ValueError("No committed import")
if target.exists(): raise ValueError("Use a fresh deliverable directory")
target.mkdir(parents=True)
shutil.copy2(prepared / "manifest.json", target / "downloads.json")
report = json.loads((prepared / "report.json").read_text(encoding="utf-8"))
rows = [json.loads(line) for line in (prepared / "candidates.jsonl").read_text(encoding="utf-8").splitlines()]
codes = {r["gtin14"] for r in rows}
if len(codes) != result["unique_codes_in_database"]: raise ValueError("Export/database count mismatch")
sets = {source: {r["gtin14"] for r in rows if r["source"] == source} for source in report["sources"]}
report["cross_source_overlaps"] = {a + " / " + b: len(sets[a] & sets[b]) for a in sets for b in sets if a < b and sets[a] & sets[b]}
reject_counts = Counter()
reject_codes = defaultdict(set)
with ExitStack() as stack:
    writers = {}
    for beauty, name in [(False, "Rechazados_detalle.csv"), (True, "Rechazados_OBF_ODbL.csv")]:
        output = stack.enter_context((target / name).open("w", encoding="utf-8-sig", newline=""))
        writers[beauty] = csv.writer(output)
        writers[beauty].writerow(["Fuente", "Codigo_original", "Nombre_original", "Motivo", "Ocurrencias", "Referencia"])
    with (prepared / "rejected.jsonl").open(encoding="utf-8") as raw:
        for line in raw:
            row = json.loads(line)
            reject_counts[row["reason"]] += row["occurrences"]
            if row["code"]: reject_codes[row["reason"]].add(row["code"])
            values = [row[k] for k in ["source", "code", "name", "reason", "occurrences", "source_ref"]]
            writers[row["source"] == "openbeauty"].writerow(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")) else v for v in values])
report["rejections_by_reason"] = {reason: {"rows": count, "distinct_nonempty_raw_codes": len(reject_codes[reason])} for reason, count in reject_counts.items()}
for beauty, filename in [(False, "Catalogo_Juncos_validos.xlsx"), (True, "Open_Beauty_Facts_ODbL.xlsx")]:
    shutil.copy2(prepared / filename, target / filename)
    book = load_workbook(target / filename)
    own = [r for r in rows if (r["source"] == "openbeauty") == beauty]
    tab = book.create_sheet("Indice de codigos unicos", 0)
    tab.append(["GTIN14", "Fuentes", "Descripciones distintas", "Estado"])
    grouped = defaultdict(list)
    for r in own: grouped[r["gtin14"]].append(r)
    for code, group in sorted(grouped.items()):
        count = len({r["fingerprint"] for r in group})
        tab.append([code, ", ".join(sorted({r["source"] for r in group})), count, "Revisar coincidencias" if count > 1 else "Pendiente de cotejo fisico"])
    tab.freeze_panes = "A2"
    tab.auto_filter.ref = tab.dimensions
    for column, width in [("A", 22), ("B", 40), ("C", 25), ("D", 35)]: tab.column_dimensions[column].width = width
    book.save(target / filename)
report["postgres"] = result
report["status"] = "IMPORTED_TO_LOCAL_POSTGRES_PRODUCTION_DISABLED"
(target / "Informe_completo.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
lines = ["# Catálogo externo Farmacia Juncos — 23/10/2026", "", f"**{len(codes):,} códigos GTIN únicos importados en PostgreSQL local independiente.** Producción desactivada; cero códigos cotejados físicamente.", "", "## Cantidades por fuente", "", "| Fuente | Filas leídas | Aceptadas (incluye repetidas) | Repetidas idénticas | Códigos únicos | Descripciones guardadas |", "|---|---:|---:|---:|---:|---:|"]
for source, s in report["sources"].items():
    lines.append(f"| {source} | {s['rows_read']:,} | {s.get('accepted_rows',0):,} | {s.get('duplicate_rows',0):,} | {s['unique_codes']:,} | {s['stored_variants']:,} |")
lines += ["", "SEPA publica filas de precios por sucursal: millones de filas NO equivalen a millones de productos. No se importan precios. Los códigos comunes entre fuentes se cuentan una sola vez en el total; sus evidencias y descripciones se conservan separadas. Las descripciones distintas no se fusionan automáticamente.", "", "## Rechazos y exclusiones", "", "| Motivo | Filas | Códigos distintos no vacíos |", "|---|---:|---:|"]
for reason, info in report["rejections_by_reason"].items(): lines.append(f"| {reason} | {info['rows']:,} | {info['distinct_nonempty_raw_codes']:,} |")
lines += ["", "`Rechazados_detalle.csv` identifica códigos, nombres, motivo, origen y ocurrencias. Los códigos pueden aparecer en varias razones/fuentes; no sumar sus cantidades como códigos únicos. Fuera de alcance y país no documentado no implican EAN inválido.", "", "## Calidad y fallos", ""]
for failure in report["failures"]: lines.append("- " + json.dumps(failure, ensure_ascii=False))
lines += ["", f"Códigos con varias descripciones: {report['codes_with_multiple_descriptions']:,}. No todos son errores: hay abreviaturas por comercio y variantes; requieren revisión.", "", "El ZIP vacío forma parte del snapshot oficial descargado (comercio 36); no contiene registros recuperables y se excluyó. El resto se procesó. Las búsquedas web de algunas páginas fallaron por 403/timeout, pero las descargas directas oficiales de ambos datasets finalizaron y tienen SHA-256 en downloads.json.", "", "La clasificación se basa en reglas conservadoras, no en una certificación sanitaria. Se excluyen medicamentos reconocidos, códigos internos, productos sin identificar y candidatos no clasificados. No se puede garantizar comercialización actual ni identidad física. SEPA acredita publicación en comercios argentinos; la etiqueta Argentina de Open Beauty Facts es colaborativa. No se infiere país por el prefijo del EAN.", "", "## Licencias y archivos", "", "- Catalogo_Juncos_validos.xlsx: SEPA CC BY 4.0 y aportes Excel/V6 identificados por separado. Derechos de los catálogos de terceros aportados: revisar antes de redistribuir.", "- Open_Beauty_Facts_ODbL.xlsx: datos de Open Beauty Facts contributors bajo ODbL 1.0; exportación independiente, sin mezclar campos con la base privada. Atribución y obligaciones de compartir bases derivadas según licencia.", "- Ambos libros incluyen índice único de códigos, evidencias completas y cantidades por marca, categoría y fuente.", "- Categoría original vacía se mantiene vacía. scope_group es una clasificación orientativa explícita, no un atributo inventado del fabricante.", "", "Fuentes: https://datos.produccion.gob.ar/dataset/sepa-precios ; https://static.openbeautyfacts.org/data/en.openbeautyfacts.org.products.csv.gz ; https://opendatacommons.org/licenses/odbl/1-0/", "", "## Seguridad", "", "Base catalogo_juncos_external, esquemas reference_catalog y beauty_odbl. Importación transaccional e idempotente, respaldo antes de actualizar una base existente. Los candidatos retirados por nuevas reglas quedan inactivos, no se borran los lotes de auditoría. Los controles de cantidades comerciales antes/después no cambiaron. No se alteraron datos ni configuración de producción. API de lectura desactivada por defecto y bloqueada en NODE_ENV=production. No hay alta automática de productos ni actualización de precios o stock."]
(target / "INFORME.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in target.iterdir() if p.is_file()}
(target / "SHA256.json").write_text(json.dumps(hashes, indent=2), encoding="utf-8")
print(json.dumps({"folder": str(target), "unique_codes": len(codes), "overlaps": report["cross_source_overlaps"], "rejections": report["rejections_by_reason"]}, indent=2))
