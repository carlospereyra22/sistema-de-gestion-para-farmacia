"""Export the actually installed local catalogue, with provenance and exclusions."""
import argparse
import csv
import json
from collections import Counter
from pathlib import Path
import psycopg
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from install import local_credentials, DB


def safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


def export(prepared, output):
    output.mkdir(parents=True, exist_ok=True)
    with psycopg.connect(**local_credentials(), dbname=DB, options='-c default_transaction_read_only=on') as conn:
        reference = [r[0] for r in conn.execute('SELECT payload FROM reference_catalog.product WHERE eligible ORDER BY gtin14')]
        beauty = [r[0] for r in conn.execute('SELECT payload FROM beauty_odbl.product WHERE eligible ORDER BY gtin14')]
        evidence = [r[0] for r in conn.execute('SELECT payload FROM reference_catalog.candidate WHERE eligible ORDER BY gtin14,source')]
    inventory = [json.loads(line) for line in (prepared / 'inventory.jsonl').read_text(encoding='utf-8').splitlines()]
    report = json.loads((prepared / 'report.json').read_text(encoding='utf-8'))
    images = json.loads((prepared / 'image-report.json').read_text(encoding='utf-8'))
    result = json.loads((prepared / 'postgres-result.json').read_text(encoding='utf-8'))
    actual_codes = {r['gtin14'] for r in reference + beauty}
    if len(actual_codes) != result['unique_codes_in_database']:
        raise ValueError('Database differs from import report')
    workbook = Workbook()
    workbook.remove(workbook.active)
    def sheet(title, headers, rows):
        ws = workbook.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append([safe(v) for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='087F73')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(65, max(16, len(str(column[0].value)) + 4))
    headers = ['Código', 'GTIN14', 'Nombre', 'Marca', 'Categoría normalizada', 'Categoría fuente', 'Presentación', 'Fuente elegida', 'Todas las fuentes', 'URL', 'Verificación', 'Revisión pendiente', 'Imagen fuente', 'Imagen disponible']
    def product_row(r):
        return [r['code'],r['gtin14'],r['name'],r['brand'],r['scope_group'],r['category'],r['presentation'],r['source'],' | '.join(r.get('sources', [])),r['source_ref'],r['verification'],r.get('review_reason', ''),r.get('image', {}).get('url', ''),bool(r.get('image'))]
    sheet('Productos referencia', headers, map(product_row, reference))
    sheet('Open Beauty Facts ODbL', headers, map(product_row, beauty))
    sheet('Evidencias originales', ['Código','Nombre original','Marca original','Categoría fuente','Fuente','URL'], ([r['code'],r['name'],r['brand'],r['category'],r['source'],r['source_ref']] for r in evidence))
    sheet('Inventario tiendas', ['Código','Nombre','Marca','Categoría fuente','Fuente','Estado','Motivo','URL'], ([r['code'],r['name'],r['brand'],r['category'],r['source'],r['status'],r['reason'],r['source_ref']] for r in inventory))
    sheet('Errores imágenes', ['GTIN14','URL','Error'], ([r['gtin14'],r['url'],r['error']] for r in images['failures']))
    quarantine_path = prepared / 'baseline-quarantine.jsonl'
    if quarantine_path.exists():
        quarantined = [json.loads(line) for line in quarantine_path.read_text(encoding='utf-8').splitlines()]
        sheet('Revisión base anterior', ['Código','Nombre','Fuente','Estado','Motivo'], ([r['code'],r['name'],r['source'],r['audit_status'],r['audit_reason']] for r in quarantined))
    for field, title in [('brand','Por marca'),('scope_group','Por categoría'),('source','Por fuente elegida')]:
        sheet(title, [title,'Códigos'], Counter(r[field] for r in reference).most_common())
    sheet('Fuentes evidencia', ['Fuente','Filas originales','Códigos únicos'], result['by_source'])
    workbook.save(output / 'CATALOGO_AMPLIADO.xlsx')
    with (output / 'PRODUCTOS_REFERENCIA.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows([[safe(v) for v in product_row(r)] for r in reference])
    summary = {'unique_codes_in_database': len(actual_codes), 'reference_codes': len(reference), 'openbeauty_codes_separate': len(beauty),
               'new_codes': report['new_codes'], 'with_cached_image': sum(bool(r.get('image')) for r in reference),
               'by_source': result['by_source'], 'downloaded_dispositions': report['disposition'], 'coverage': report['coverage'],
               'retail_unique_codes': report['retail_unique_codes'], 'retail_duplicate_rows': report['retail_duplicate_rows'],
               'retail_codes_already_in_baseline': report['retail_codes_already_in_baseline'],
               'invalid_code_reasons': report['invalid_code_reasons'],
               'commercial_counts_unchanged': result['commercial_counts_unchanged'], 'production_enabled': False}
    (output / 'RESULTADO.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    text = f"""# Catálogo externo: ampliación de tiendas públicas

Datos consultados en PostgreSQL local después de importar: **{len(actual_codes)} códigos únicos**.
Referencia: {len(reference)}; Open Beauty Facts: {len(beauty)} en colección ODbL separada (pueden solaparse).
Códigos nuevos respecto a la base anterior: **{report['new_codes']}**.
Fotos descargadas y disponibles: **{summary['with_cached_image']}**; intentos fallidos: {len(images['failures'])}.

Los códigos se validaron por formato y dígito de control. No se concedió verificación física.
SKU interno, código ausente o inválido: excluido. Medicamentos y registros fuera del alcance no se incorporan como productos no medicinales.
Variantes descargadas: {report['downloaded_variants']}. Admitidas: {report['disposition']['candidate']} filas, equivalentes a {report['retail_unique_codes']} códigos únicos de las tres tiendas.
Duplicados entre esas filas: {report['retail_duplicate_rows']}. Códigos ya presentes en la base anterior: {report['retail_codes_already_in_baseline']}.
Rechazadas por código: {report['disposition']['invalid_code']}; excluidas por alcance: {report['disposition']['discard']}; pendientes de revisión: {report['disposition']['review']}.
Las discrepancias entre nombres, marcas o presentación conservan advertencias. Una ficha canónica por GTIN conserva las evidencias originales.
Farmacorp es una fuente de Bolivia: comercialización argentina pendiente salvo coincidencia exacta con fuentes argentinas.

## Cobertura y límites

RESULTADO.json detalla totales declarados, páginas y categorías recorridas y URLs pendientes. Los totales de la web pueden variar mientras se descarga.
El inventario incluye registros descartados y pendientes de revisión: no se cuentan como importados.
Farmar corresponde a la sucursal pública 136; no acredita inventario privado ni otras sucursales. Su sitemap de productos estaba vacío; se usó el listado público paginado.
Las 113 categorías coinciden con sus totales declarados. 805 fichas del listado general no aparecieron en esas categorías: se conserva esa ausencia, sin atribuirles una categoría de tienda inventada.
Farmalife se contrastó contra sus siete sitemaps de productos. Las solicitudes fallidas se conservan en el reporte de cobertura.
Farmacorp se paginó hasta la respuesta vacía. Los códigos internos que no cumplen GTIN no se convierten en EAN.

## Imágenes y uso

Una imagen por referencia cuando pudo asociarse al código y descargarse. JPEG optimizado hasta 60 KB, con URL original y huella SHA-256.
No se atribuye licencia abierta a las fotos de las tiendas; permanecen sujetas a las condiciones de sus titulares.
Se muestran en la vista previa local y al seleccionar una ficha completan la foto del borrador, sin reemplazar una foto manual.
No se incorporan precios ni existencias de las tiendas. Los datos comerciales existentes permanecen intactos.
Integración productiva: **desactivada**.

## Archivos

- CATALOGO_AMPLIADO.xlsx: referencia, ODbL separado, evidencias, inventario completo descargado, marcas, categorías y errores de imágenes.
- PRODUCTOS_REFERENCIA.csv: una fila por GTIN de referencia.
- RESULTADO.json: cantidades y cobertura comprobadas.
"""
    (output / 'INFORME.md').write_text(text, encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k != 'coverage'}, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('prepared', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    export(args.prepared, args.output)
