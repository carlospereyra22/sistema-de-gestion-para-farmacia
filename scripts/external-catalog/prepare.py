"""Adapted SEPA importer: streams actual sources, no fabricated fields or commercial writes.
Run with uv run --python 3.12 --with openpyxl==3.1.5 ...
"""
import argparse
import csv
import gzip
import hashlib
import io
import json
import pathlib
import tempfile
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from rules import text, gtin_error, classify

FIELDS = ["code", "gtin14", "name", "brand", "category", "presentation", "source", "source_ref", "license", "verification", "scope_group", "argentina_evidence", "fingerprint", "occurrences"]

class Importer:
    def __init__(self, out):
        self.out = out
        out.mkdir(parents=True, exist_ok=True)
        self.entries = {}
        self.stats = defaultdict(Counter)
        self.failures = []
        self.rejected = (out / "rejected.jsonl").open("w", encoding="utf-8")

    def row(self, source, code, name, brand="", category="", presentation="", ref="", license="", evidence="", flag=None, count=1):
        stats = self.stats[source]
        stats["rows_read"] += count
        code, name, brand, category, presentation = map(text, [code, name, brand, category, presentation])
        reason = gtin_error(code)
        if not reason and flag is not None and flag.lower() not in {"1", "true", "si", "sí"}:
            reason = "not_explicitly_ean"
        group, scope_error = classify(name, category)
        reason = reason or ("missing_name" if not name else None) or scope_error
        if reason:
            stats["rejected"] += count
            stats[reason] += count
            self.rejected.write(json.dumps({"source": source, "code": code, "name": name, "reason": reason, "occurrences": count, "source_ref": ref}, ensure_ascii=False) + "\n")
            return
        if source == "openbeauty" and evidence != "en:argentina":
            stats["rejected"] += count
            stats["not_argentina"] += count
            self.rejected.write(json.dumps({"source": source, "code": code, "name": name, "reason": "not_argentina", "occurrences": count, "source_ref": ref}, ensure_ascii=False) + "\n")
            return
        gtin14 = code.zfill(14)
        # Preserve complete, independent descriptions. No picking a longest name and mixing brands.
        fingerprint = hashlib.sha256(json.dumps([gtin14, name, brand, category, presentation], ensure_ascii=False).encode()).hexdigest()
        key = source, fingerprint
        stats["accepted_rows"] += count
        if key in self.entries:
            stats["duplicate_rows"] += count
            self.entries[key]["occurrences"] += count
        else:
            stats["duplicate_rows"] += count - 1
            self.entries[key] = dict(zip(FIELDS, [code, gtin14, name, brand, category, presentation, source, ref, license, "CHECKSUM_VALID_NOT_PHYSICALLY_VERIFIED", group, evidence, fingerprint, count]))

    def excel(self, path):
        book = load_workbook(path, read_only=True, data_only=False)
        sheet = book["Productos con EAN"]
        rows = sheet.iter_rows()
        next(rows)
        for row in rows:
            values = [c.value for c in row[:8]]
            if not any(v is not None for v in values):
                continue
            if any(c.data_type == "f" for c in row[:8]):
                self.failures.append({"source": "excel270", "row": row[0].row, "error": "formula_in_product_row"})
                continue
            code, name, brand, category, presentation, ref, _, _ = values
            self.row("excel270", code, name, brand, category, presentation, text(ref) + f" [Excel row {row[0].row}]", "USER_SUPPLIED_REVIEW_RIGHTS", "user_supplied_argentina")
        book.close()

    def v6(self, path):
        with zipfile.ZipFile(path) as archive:
            with archive.open("Catalogo_Juncos_V6/data/productos_con_ean_publicado_v4.csv") as raw:
                for r in csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")):
                    self.row("v6", r["gtin"], r["name"], r["brand"], r["category"], r["presentation"], r["source_url"], "USER_SUPPLIED_REVIEW_RIGHTS", "user_supplied_argentina")
            with archive.open("Catalogo_Juncos_V6/data/productos_farmacias_sin_ean.csv") as raw:
                for r in csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")):
                    self.row("v6_without_ean", "", r.get("name", r.get("nombre", "")), ref=path.name)

    def sepa(self, archive, ref, depth=0):
        if depth > 3:
            raise ValueError("Nested ZIP limit exceeded")
        for entry in archive.infolist():
            if entry.is_dir():
                continue
            location = ref + "!" + entry.filename
            try:
                if entry.file_size > 12_000_000_000:
                    raise ValueError("Uncompressed member exceeds safety limit")
                if entry.filename.lower().endswith(".zip"):
                    if not entry.file_size:
                        raise ValueError("Empty ZIP supplied by SEPA")
                    with tempfile.SpooledTemporaryFile(max_size=32_000_000) as temp, archive.open(entry) as raw:
                        while chunk := raw.read(1024 * 1024):
                            temp.write(chunk)
                        temp.seek(0)
                        with zipfile.ZipFile(temp) as nested:
                            self.sepa(nested, location, depth + 1)
                elif entry.filename.lower().endswith("productos.csv"):
                    records = Counter()
                    with archive.open(entry) as raw:
                        reader = csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline=""), delimiter="|")
                        required = {"id_producto", "productos_ean", "productos_descripcion"}
                        if not required.issubset(reader.fieldnames or []):
                            raise ValueError("Unrecognized SEPA columns")
                        for r in reader:
                            code = text(r.get("id_producto"))
                            name = text(r.get("productos_descripcion"))
                            # Quantity and unit are actual source fields, not parsed/inferred from a name.
                            presentation = " ".join(filter(None, [text(r.get("productos_cantidad_presentacion")), text(r.get("productos_unidad_medida_presentacion"))]))
                            records[(code, name, text(r.get("productos_marca")), presentation, text(r.get("productos_ean")))] += 1
                    for (code, name, brand, presentation, flag), count in records.items():
                        self.row("sepa", code, name, brand, presentation=presentation, ref=location, license="CC-BY-4.0", evidence="SEPA Argentina publication", flag=flag, count=count)
                    print("SEPA processed", entry.filename, "distinct rows", len(records), flush=True)
            except Exception as error:
                self.failures.append({"source": "sepa", "file": location, "error": str(error)})
                print("SEPA member failed", entry.filename, str(error), flush=True)

    def beauty(self, path):
        csv.field_size_limit(16_000_000)
        with gzip.open(path, "rt", encoding="utf-8-sig", newline="") as raw:
            reader = csv.DictReader(raw, delimiter="\t")
            if not {"code", "product_name", "countries_tags"}.issubset(reader.fieldnames or []):
                raise ValueError("Invalid Open Beauty Facts columns")
            for r in reader:
                tags = text(r.get("countries_tags")).split(",")
                self.row("openbeauty", r["code"], r.get("product_name_es") or r["product_name"], r.get("brands"), r.get("categories"), r.get("quantity"), r.get("url"), "ODbL-1.0", "en:argentina" if "en:argentina" in tags else text(r.get("countries_tags")))

    def supplier(self, path, source, license):
        if not source or not license:
            raise ValueError("Supplier import requires source and explicit authorization reference")
        with path.open(encoding="utf-8-sig", newline="") as raw:
            reader = csv.DictReader(raw)
            if not {"code", "name", "brand", "category", "presentation"}.issubset(reader.fieldnames or []):
                raise ValueError("Expected columns: code,name,brand,category,presentation,source_ref")
            for r in reader:
                self.row("supplier:" + source, r["code"], r["name"], r["brand"], r["category"], r["presentation"], r.get("source_ref", path.name), license, "authorized_supplier_argentina")

    def finish(self):
        self.rejected.close()
        entries = list(self.entries.values())
        by_code = defaultdict(set)
        for r in entries:
            by_code[r["gtin14"]].add(r["fingerprint"])
        for source, stats in self.stats.items():
            own = [r for r in entries if r["source"] == source]
            stats["unique_codes"] = len({r["gtin14"] for r in own})
            stats["stored_variants"] = len(own)
        report = {"prepared_at": datetime.now(timezone.utc).isoformat(), "status": "PREPARED_NOT_YET_IMPORTED_TO_POSTGRES", "sources": dict(self.stats), "unique_codes": len(by_code), "codes_with_multiple_descriptions": sum(len(v) > 1 for v in by_code.values()), "failures": self.failures, "physically_verified": 0}
        for column in ["brand", "category", "source", "scope_group"]:
            counts = defaultdict(set)
            for r in entries:
                counts[r[column] or "Sin informar"].add(r["gtin14"])
            report["by_" + column] = {k: len(v) for k, v in sorted(counts.items())}
        with (self.out / "candidates.jsonl").open("w", encoding="utf-8") as f:
            for r in sorted(entries, key=lambda x: (x["source"], x["gtin14"], x["fingerprint"])):
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        for beauty in [False, True]:
            own = [r for r in entries if (r["source"] == "openbeauty") == beauty]
            book = Workbook()
            sheet = book.active
            sheet.title = "Productos candidatos"
            sheet.append(FIELDS)
            for row_number, r in enumerate(own, start=2):
                sheet.append([r[k] for k in FIELDS])
                # Source strings must never become executable Excel formulas.
                for column in range(1, len(FIELDS) + 1):
                    cell = sheet.cell(row_number, column)
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for c in sheet[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="087F73")
            for col, width in {"A": 20, "B": 20, "C": 65, "D": 25, "E": 35, "F": 25, "G": 20, "H": 65, "I": 30, "J": 42}.items():
                sheet.column_dimensions[col].width = width
            notice = book.create_sheet("Fuentes y alcance")
            notice.append(["No hay productos cotejados físicamente; revisar coincidencias antes de activar."])
            notice.append(["Categoría se conserva de fuente. scope_group es clasificación orientativa por reglas, no dato verificado."])
            notice.append(["Variantes contradictorias del mismo código se conservan; no se mezclan campos."])
            notice.append(["Open Beauty Facts contributors — ODbL 1.0 — https://opendatacommons.org/licenses/odbl/1-0/" if beauty else "SEPA Argentina — CC BY 4.0 — https://datos.produccion.gob.ar/dataset/sepa-precios. Excel/V6: aportados por usuario; derechos de terceros pendientes de revisión."])
            for field in ["brand", "category", "source"]:
                tab = book.create_sheet("Por " + field)
                tab.append([field, "Códigos únicos (no sumar entre fuentes)"])
                values = defaultdict(set)
                for r in own:
                    values[r[field] or "Sin informar"].add(r["gtin14"])
                for label, codes in sorted(values.items()):
                    tab.append([label, len(codes)])
                    tab.cell(tab.max_row, 1).data_type = "s"
            book.save(self.out / ("Open_Beauty_Facts_ODbL.xlsx" if beauty else "Catalogo_Juncos_validos.xlsx"))
        (self.out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: report[k] for k in ["status", "sources", "unique_codes", "codes_with_multiple_descriptions", "failures"]}, ensure_ascii=True, indent=2), flush=True)

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--input-dir", type=pathlib.Path, required=True)
    p.add_argument("--out", type=pathlib.Path, required=True)
    p.add_argument("--excel", type=pathlib.Path)
    p.add_argument("--v6", type=pathlib.Path)
    p.add_argument("--supplier-csv", type=pathlib.Path)
    p.add_argument("--source")
    p.add_argument("--authorization")
    a = p.parse_args()
    if a.out.exists():
        raise ValueError("Use a new output directory; never overwrite a previous run")
    importer = Importer(a.out)
    for source in ["excel270", "v6", "sepa", "openbeauty", "supplier"]:
        try:
            if source == "excel270" and a.excel: importer.excel(a.excel)
            elif source == "v6" and a.v6: importer.v6(a.v6)
            elif source == "sepa" and (a.input_dir / "sepa.zip").exists():
                with zipfile.ZipFile(a.input_dir / "sepa.zip") as archive: importer.sepa(archive, "https://datos.produccion.gob.ar/dataset/sepa-precios [sepa.zip]")
            elif source == "openbeauty" and (a.input_dir / "openbeauty.csv.gz").exists(): importer.beauty(a.input_dir / "openbeauty.csv.gz")
            elif source == "supplier" and a.supplier_csv: importer.supplier(a.supplier_csv, a.source, a.authorization)
        except Exception as error:
            importer.failures.append({"source": source, "error": str(error)})
    importer.finish()

if __name__ == "__main__": main()
