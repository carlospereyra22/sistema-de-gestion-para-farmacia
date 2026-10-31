"""Verify deliverables against actual rows, not estimates, and export review samples."""
import argparse
import csv
import json
import pathlib
import sqlite3
from collections import Counter, defaultdict
from contextlib import closing
from openpyxl import load_workbook
from rules import gtin_error


def validate(folder):
    report = json.loads((folder / "report.json").read_text(encoding="utf-8"))
    summary = report["directory_summary"]
    groups = {}
    for filename in ["CATALOGO_JUNCOS_FARMACIA_COMPLETO.csv", "CATALOGO_JUNCOS_REVISION.csv", "OPEN_BEAUTY_FACTS_ODbL.csv"]:
        with (folder / filename).open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        assert all(not gtin_error(r["codigo_barras"]) and r["gtin14"] == r["codigo_barras"].zfill(14) for r in rows)
        assert len({r["gtin14"] for r in rows}) == len(rows)
        groups[filename] = rows
    main, review, beauty = groups.values()
    assert len(main) == summary["main_classified"]
    assert len(review) == summary["review_rows"]
    assert len(beauty) == summary["openbeauty_codes"]
    assert not ({r["gtin14"] for r in main} & {r["gtin14"] for r in review})
    assert dict(Counter(r["categoria"] for r in main)) == summary["by_category"]
    with (folder / "CATALOGO_JUNCOS_DESCARTADOS.csv").open(encoding="utf-8-sig", newline="") as f:
        discarded_count = 0
        occurrences = 0
        for row in csv.DictReader(f):
            discarded_count += 1
            occurrences += int(row["ocurrencias"])
    assert occurrences == summary["discarded_rows"]
    candidates = [json.loads(line) for line in (folder / "candidates.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len({r["gtin14"] for r in candidates}) == summary["combined_codes"]
    assert len({(r["source"], r["fingerprint"]) for r in candidates}) == len(candidates)
    assert all(not gtin_error(r["code"]) for r in candidates)
    assert all(not ({"price", "stock", "promotion"} & r.keys()) for r in candidates)
    with closing(sqlite3.connect(folder / "CATALOGO_JUNCOS_FARMACIA.sqlite")) as db:
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert db.execute("SELECT count(*) FROM productos").fetchone()[0] == len(main) + len(review)
        assert db.execute("SELECT count(*) FROM evidencias WHERE source='openbeauty'").fetchone()[0] == 0
        for row in main:
            assert db.execute("SELECT nombre FROM productos WHERE codigo_barras=?", (row["codigo_barras"],)).fetchone()[0] == row["nombre_producto"]
        searches = db.execute("SELECT codigo_barras,nombre FROM productos WHERE nombre LIKE ? LIMIT 5", ("%PANTENE%",)).fetchall()
        assert searches
    workbook = load_workbook(folder / "CATALOGO_JUNCOS_RESUMEN.xlsx", read_only=True)
    try:
        for sheet, count in [("Productos clasificados", len(main)), ("Revision", len(review))]:
            assert sum(1 for _ in workbook[sheet].rows) - 1 == count
        first_code = next(workbook["Productos clasificados"].iter_rows(min_row=2, max_row=2))[0]
        assert first_code.data_type == "s"
    finally:
        workbook.close()
    by_category = defaultdict(list)
    for row in main:
        by_category[row["categoria"]].append(row)
    with (folder / "MUESTRAS_POR_CATEGORIA.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(main[0]))
        writer.writeheader()
        for rows in by_category.values():
            writer.writerows(rows[::max(1, len(rows) // 5)][:5])
    result = dict(status="PASSED", approved_csv=len(main), review_csv=len(review), discarded_distinct_source_rows=discarded_count, discarded_occurrences=occurrences, evidence_rows=len(candidates), codes=len({r["gtin14"] for r in candidates}), sqlite_integrity="ok", barcode_lookup="all_classified_codes_match", name_search=searches, workbook_counts="match", formulas_not_executed=True, production_enabled=False)
    (folder / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("folder", type=pathlib.Path)
    validate(parser.parse_args().folder)
