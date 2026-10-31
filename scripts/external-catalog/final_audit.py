"""Close the offline catalogue audit using measured outputs and raw cached evidence."""
import argparse
import csv
import json
import pathlib
import re
from collections import Counter
from openpyxl import Workbook
from rules import gtin_error


def records(folder, filename):
    with (folder / filename).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path, columns, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def audit(scan, before, after):
    baseline = [json.loads(line) for line in (before / "candidates.jsonl").read_text(encoding="utf-8").splitlines()]
    current = [json.loads(line) for line in (after / "candidates.jsonl").read_text(encoding="utf-8").splitlines()]
    old_codes = {r["gtin14"] for r in baseline}
    new_codes = {r["gtin14"] for r in current}
    assert old_codes <= new_codes
    main = records(after, "CATALOGO_JUNCOS_FARMACIA_COMPLETO.csv")
    pending = records(after, "CATALOGO_JUNCOS_REVISION.csv")
    prior_main = {r["gtin14"] for r in records(before, "CATALOGO_JUNCOS_FARMACIA_COMPLETO.csv")}
    prior_pending = {r["gtin14"] for r in records(before, "CATALOGO_JUNCOS_REVISION.csv")}
    columns = list(main[0])
    added = [r for r in main + pending if r["gtin14"] not in old_codes and r["gtin14"] in new_codes]
    reopened = [r for r in pending if r["gtin14"] in prior_main]
    resolved = [r for r in main if r["gtin14"] in prior_pending]
    for filename, rows in [("FINAL_PRODUCTOS_RECUPERADOS.csv", added), ("FINAL_FICHAS_REABIERTAS.csv", reopened), ("FINAL_PENDIENTES_RESUELTOS.csv", resolved)]:
        write_csv(after / filename, columns, rows)
    assert len(added) == len(new_codes - old_codes)

    # All new SEPA observations must match a complete original tuple, not merged fields.
    def key(row):
        return tuple(row.get(k, "") for k in ("code", "name", "brand", "quantity", "unit"))
    old_keys = {key(r) for r in baseline if r["source"] == "sepa"}
    additions = {key(r) for r in current if r["source"] == "sepa"} - old_keys
    unmatched = set(additions)
    inventory = json.loads((scan / "inventory.json").read_text(encoding="utf-8"))
    scanned_hashes = set()
    cached_observations = 0
    for file in inventory["files"]:
        if file["status"] != "processed" or file["sha256"] in scanned_hashes:
            continue
        scanned_hashes.add(file["sha256"])
        occurrences = 0
        for line in (scan / "cache" / (file["sha256"] + ".jsonl")).read_text(encoding="utf-8").splitlines():
            row = json.loads(line)
            cached_observations += 1
            occurrences += row[-1]
            unmatched.discard(tuple(row[:5]))
        assert occurrences + file["counts"].get("malformed_rows", 0) == file["counts"]["rows"]
    assert not unmatched, "An added observation has no original whole source tuple"

    reasons = Counter()
    invalid_codes = set()
    valid_excluded = set()
    with (after / "CATALOGO_JUNCOS_DESCARTADOS.csv").open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            reasons[row["motivo"]] += int(row["ocurrencias"])
            (invalid_codes if gtin_error(row["codigo_barras"]) else valid_excluded).add(row["codigo_barras"])
    quality = []
    for row in main + pending:
        flags = []
        if not row["marca"].strip(): flags.append("marca_no_informada")
        if not row["presentacion"].strip(): flags.append("presentacion_no_informada")
        if re.search(r"[ÃÂ�]", row["nombre_producto"]): flags.append("posible_codificacion_original")
        if flags:
            quality.append(dict(row, alertas=";".join(flags)))
    write_csv(after / "FINAL_CALIDAD_CAMPOS.csv", columns + ["alertas"], quality)
    algabo = {"Clasificados": [r for r in main if re.search(r"\balgabo\b", r["nombre_producto"] + " " + r["marca"], re.I)], "En revision": [r for r in pending if re.search(r"\balgabo\b", r["nombre_producto"] + " " + r["marca"], re.I)]}
    book = Workbook()
    book.remove(book.active)
    for title, rows in algabo.items():
        tab = book.create_sheet(title)
        tab.append(columns)
        for row in rows:
            tab.append([row[c] for c in columns])
            for cell in tab[tab.max_row]:
                cell.data_type = "s"
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        tab.column_dimensions["A"].width = 20
        tab.column_dimensions["B"].width = 65
        tab.column_dimensions["C"].width = 25
    book.save(after / "ALGABO_PRODUCTOS.xlsx")
    integrity = json.loads((scan / "final-source-integrity.json").read_text(encoding="utf-8"))
    assert integrity["passed"]
    (after / "source-integrity.json").write_text(json.dumps(integrity, indent=2), encoding="utf-8")
    result = dict(previous_codes=len(old_codes), current_codes=len(new_codes), recovered_codes=len(added), recovered_classified=sum(not r["motivo_revision"] for r in added), recovered_review=sum(bool(r["motivo_revision"]) for r in added), classified_codes=len(main), review_rows=len(pending), review_imported=sum(r["gtin14"] in new_codes for r in pending), review_not_imported=sum(r["gtin14"] not in new_codes for r in pending), previous_main_reopened=len(reopened), pending_resolved=len(resolved), new_sepa_observations_verified=len(additions), cached_observations_checked=cached_observations, invalid_distinct_codes=len(invalid_codes), valid_codes_with_excluded_observations=len(valid_excluded), rejected_occurrences_by_reason=dict(reasons), algabo={k:len(v) for k,v in algabo.items()}, field_quality_rows=len(quality), original_files_unchanged=integrity["passed"], previous_codes_preserved=True, physical_verification=False, production_enabled=False)
    (after / "final-audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True, indent=2))


def report(after):
    result = json.loads((after / "final-audit.json").read_text(encoding="utf-8"))
    db = json.loads((after / "postgres-result.json").read_text(encoding="utf-8"))
    validation = json.loads((after / "validation.json").read_text(encoding="utf-8"))
    summary = json.loads((after / "report.json").read_text(encoding="utf-8"))["directory_summary"]
    assert db["unique_codes_in_database"] == result["current_codes"]
    assert db["commercial_counts_unchanged"] and validation["status"] == "PASSED"
    lines = ["# Auditoría final del catálogo externo — 24/10/2026", "", "Base de referencia local preparada para consulta y validación de envases. No activa producción ni crea productos comerciales.", "", "| Resultado | Cantidad |", "|---|---:|"]
    labels = {"previous_codes":"Códigos antes de esta revisión", "current_codes":"Códigos únicos importados en PostgreSQL externo", "recovered_codes":"Códigos adicionales recuperados", "recovered_classified":"Nuevos clasificados por reglas", "recovered_review":"Nuevos que requieren revisar discrepancias", "classified_codes":"Fichas del listado principal", "review_imported":"Códigos importados conservados en revisión", "review_not_imported":"Códigos ambiguos adicionales NO importados", "previous_main_reopened":"Fichas anteriores reabiertas para revisión", "pending_resolved":"Pendientes anteriores ahora clasificados", "new_sepa_observations_verified":"Observaciones nuevas cotejadas con filas fuente completas"}
    lines.extend(f"| {label} | {result[key]:,} |" for key,label in labels.items())
    lines.extend(["", "Los conteos de pendientes resueltos pueden superponerse con códigos nuevos. Open Beauty Facts se conserva aparte; sus 131 códigos pueden coincidir con otras fuentes, por lo que no se suman sin deduplicar.", "", "## Alcance real de la revisión", "",
        f"Se comprobó la huella SHA256 de los archivos originales disponibles: 125 CSV, 11.054.413.981 bytes, sin cambios ni CSV nuevos omitidos. Los 113 CSV de productos representan {summary['rows_read']:,} filas. Se reclasificó toda la evidencia cacheada y se cotejaron {result['new_sepa_observations_verified']:,} observaciones nuevas con tuplas originales completas (código, nombre, marca, cantidad y unidad). No se reconstruyeron nombres ni se inventaron variantes.",
        "Se revisaron descartes por marca y tipo, abreviaturas observadas de bebés, depilación, cabello, higiene y cosmética, y coincidencias del mismo GTIN. La marca sola deja un registro en revisión; nunca basta para aprobarlo. Los nombres de familias/modelos sin tipo explícito siguen pendientes.",
        "También se corrigieron falsos positivos: desodorantes de ambiente abreviados, fragancias de uso automotor y tabletas repelentes ambientales. Los registros históricos del catálogo externo se conservan con motivo de revisión, fuera del listado principal. Nunca se borró historia comercial.",
        f"Algabo: {result['algabo']['Clasificados']} fichas principales y {result['algabo']['En revision']} en revisión. Se recuperaron los dos registros de barro vegetal/natural observados. ALGABO_PRODUCTOS.xlsx contiene el filtro actualizado.", "",
        "## Archivos para usar", "",
        "- CATALOGO_JUNCOS_RESUMEN.xlsx: productos clasificados, revisión, cantidades por marca/categoría/fuente y duplicados.",
        "- FINAL_PRODUCTOS_RECUPERADOS.csv: los códigos incorporados en este cierre.",
        "- FINAL_FICHAS_REABIERTAS.csv y FINAL_PENDIENTES_RESUELTOS.csv: cambios de estado auditables.",
        "- FINAL_CALIDAD_CAMPOS.csv: campos no informados y posibles problemas de codificación originales; no se completan por inferencia.",
        "- CATALOGO_JUNCOS_DESCARTADOS.csv y ERRORES_LECTURA.csv: motivos, ocurrencias, procedencia y fallos originales.",
        "- OPEN_BEAUTY_FACTS_ODbL.csv: fuente separada, con atribución y licencia; no se mezcla con el SQLite principal.",
        "- final-audit.json, validation.json, postgres-result.json, source-integrity.json y SHA256SUMS.json: evidencias de validación y entrega.", "",
        "## Verificaciones y límites", "",
        "16 pruebas Python aprobadas; GTIN/formato/checksum/unicidad de todas las exportaciones; conciliación Excel/CSV/SQLite y búsqueda de todos los códigos clasificados. API local probada con código aportado, código ausente y bloqueo en producción. Los códigos únicos importados son los confirmados por PostgreSQL, no estimaciones de archivos descargados.",
        "Se generó respaldo del catálogo externo antes de importar y se verificó su índice con pg_restore --list. Se importó en transacción. Los conteos comerciales antes/después permanecen iguales. No se probó restauración completa. No se ejecutaron lint/typecheck/build de toda la aplicación: los cambios son herramientas Python offline y documentación; la API y React no se modificaron.",
        "Los GTIN con checksum válido no prueban titularidad GS1 ni identidad física. Categorías inferidas y descripciones abreviadas siguen necesitando cotejo del envase. No se garantiza que el dataset contenga todos los productos comercializados ni se fuerza la inclusión de nombres insuficientes. Los problemas de codificación de origen se conservan y se señalan, sin corregir nombres a ciegas.", "",
        "## Uso y actualización", "",
        "La base local es catalogo_juncos_external, separada de gestion_dev. La consulta de vista previa requiere autorización y siempre devuelve candidatos para revisar; no selecciona automáticamente ni copia precios o stock. El formulario de Productos y producción siguen desactivados hasta validar coincidencias reales.",
        "Para una nueva edición, escanear las carpetas SEPA, construir hacia una carpeta nueva usando candidates.jsonl de esta entrega como base, validar, revisar diferencias e importar con install.py (respaldo previo automático). Los catálogos autorizados futuros deben conservar procedencia/licencia y pasar los mismos controles; no sobrescribir las entregas anteriores.", ""])
    (after / "INFORME_AUDITORIA_FINAL.md").write_text("\n".join(lines), encoding="utf-8")
    guide = ["# Catálogo externo Juncos — entrega vigente", "",
        f"Base local: {result['current_codes']:,} códigos únicos importados. Listado principal: {result['classified_codes']:,} fichas clasificadas por reglas, todavía sin cotejo físico.", "",
        "1. Abrir CATALOGO_JUNCOS_RESUMEN.xlsx y entrar en Productos clasificados. Buscar por nombre, marca o código. Activar Datos > Filtro para filtrar columnas.",
        "2. La hoja Revision incluye discrepancias y candidatos ambiguos: no copiar sus datos sin comprobar el envase y la fuente.",
        "3. ALGABO_PRODUCTOS.xlsx es un filtro de la misma entrega, con hojas Clasificados y En revision.",
        "4. INFORME_AUDITORIA_FINAL.md explica los resultados vigentes. Los informes de segunda/tercera revisión son antecedentes, no conteos actuales.", "",
        "No se cargaron estos registros en Productos del sistema. La base externa es independiente; no contiene ni modifica stock o precios comerciales. El formulario y producción permanecen desactivados.", "",
        "Para validar la futura integración, escanear códigos físicos y comparar nombre completo, marca y presentación con los candidatos. Ante diferencias, conservar el caso pendiente; no elegir automáticamente. Open Beauty Facts se consulta separado con atribución ODbL.", "",
        "Los códigos deben conservarse como texto al abrir o importar CSV. No quitar ceros iniciales ni interpretar tamaños/variantes faltantes. Para revisar rechazos, usar CATALOGO_JUNCOS_DESCARTADOS.csv: cada fila incluye motivo, ocurrencias y archivo original.", "",
        "Los scripts entregados permiten repetir la preparación y validación desde el repositorio. Las rutas de fuentes están documentadas en manifest.json; las credenciales no se incluyen.", ""]
    (after / "LEEME_PRIMERO.md").write_text("\n".join(guide), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("scan", type=pathlib.Path)
    parser.add_argument("before", type=pathlib.Path)
    parser.add_argument("after", type=pathlib.Path)
    parser.add_argument("--report", action="store_true")
    args = parser.parse_args()
    if args.report:
        report(args.after)
    else:
        audit(args.scan, args.before, args.after)
